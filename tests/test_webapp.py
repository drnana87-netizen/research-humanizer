"""Integration smoke tests. Run: python -m unittest discover -s tests -v"""
import base64
import io
import json
import os
import sys
import threading
import unittest
from http.server import ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch, MagicMock
import types

from docx import Document
import requests

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import server


class ResearchHumanizerWebTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        os.environ['ALLOW_HTTP_COOKIES'] = '1'
        os.environ['APP_ACCESS_PASSWORD'] = 'this-is-a-strong-private-passphrase-for-tests'
        os.environ['OPENAI_API_KEY'] = 'test-never-send-to-client'
        cls.http = ThreadingHTTPServer(('127.0.0.1', 0), server.Handler)
        cls.thread = threading.Thread(target=cls.http.serve_forever, daemon=True)
        cls.thread.start()
        cls.origin = f'http://127.0.0.1:{cls.http.server_port}'

    @classmethod
    def tearDownClass(cls):
        cls.http.shutdown()
        cls.http.server_close()

    def setUp(self):
        self.client = requests.Session()
        self.headers = {'Origin': self.origin}

    def login(self):
        resp = self.client.post(self.origin + '/api/login', json={'password': os.environ['APP_ACCESS_PASSWORD']}, headers=self.headers)
        self.assertEqual(resp.status_code, 200, resp.text)
        return resp

    def test_site_serves_and_keeps_secret_off_browser(self):
        response = self.client.get(self.origin + '/')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Unlock Research Humanizer', response.text)
        self.assertNotIn(os.environ['OPENAI_API_KEY'], response.text)
        self.assertIn("frame-ancestors 'none'", response.headers['Content-Security-Policy'])
        self.assertEqual(self.client.get(self.origin+'/api/status').status_code, 401)

    def test_authentication_session_and_logout(self):
        self.assertEqual(self.client.post(self.origin+'/api/login', json={'password':'incorrect-long-password'},headers=self.headers).status_code,401)
        resp = self.login()
        self.assertIn('HttpOnly', resp.headers['Set-Cookie'])
        self.assertIn('SameSite=Strict', resp.headers['Set-Cookie'])
        status = self.client.get(self.origin+'/api/status')
        self.assertEqual(status.status_code,200)
        self.assertEqual(status.json()['ready'],True)
        self.assertNotIn('test-never-send',status.text)
        self.assertEqual(self.client.post(self.origin+'/api/logout', json={},headers=self.headers).status_code,200)
        self.assertEqual(self.client.get(self.origin+'/api/status').status_code,401)

    def test_reject_csrf(self):
        self.login()
        response = self.client.post(self.origin + '/api/rewrite',json={'text':'Some long research passage that can be rewritten.'}, headers={'Origin':'https://evil.example'})
        self.assertEqual(response.status_code, 403)
        response = self.client.post(self.origin + '/api/rewrite',json={'text':'Some long research passage that can be rewritten.'})
        self.assertEqual(response.status_code, 403)

    def test_docx_import_and_export(self):
        self.login()
        doc = Document()
        doc.add_paragraph('The sample included 350 participants.')
        doc.add_paragraph('References')
        doc.add_paragraph('Doe, J. (2024). Example only.')
        b = io.BytesIO(); doc.save(b)
        response=self.client.post(self.origin+'/api/extract',json={'name':'review.docx','base64':base64.b64encode(b.getvalue()).decode()},headers=self.headers)
        self.assertEqual(response.status_code,200,response.text)
        self.assertIn('350 participants',response.json()['text'])
        saved=self.client.post(self.origin+'/api/export-docx',json={'text':'Literature Review\nEvidence matters.'},headers=self.headers)
        self.assertEqual(saved.status_code,200,saved.text)
        parsed=Document(io.BytesIO(saved.content))
        self.assertIn('Evidence matters.', '\n'.join(p.text for p in parsed.paragraphs))

    def test_pdf_import(self):
        import fitz
        self.login()
        pdf = fitz.open()
        page = pdf.new_page()
        page.insert_text((70, 72), 'The trial reported 48 participants and a 5% response.')
        raw = pdf.tobytes()
        response = self.client.post(self.origin+'/api/extract', json={'name':'article.pdf','base64':base64.b64encode(raw).decode()},headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn('48 participants', response.json()['text'])

    def test_plain_text_import(self):
        self.login()
        response=self.client.post(self.origin+'/api/extract',json={'name':'test.txt','base64':base64.b64encode('Arabic عربي and English'.encode()).decode()},headers=self.headers)
        self.assertEqual(response.status_code,200,response.text)
        self.assertIn('عربي',response.json()['text'])

    def test_api_rewrite_uses_private_key_and_store_false(self):
        self.login()
        mock_api = types.ModuleType('openai')
        ctor = MagicMock()
        mock_api.OpenAI = ctor
        with patch.dict(sys.modules, {'openai': mock_api}):
            ctor.return_value.responses.create.return_value.output_text='Edited evidence (Smith, 2022) with 42 cases.'
            text='The evidence (Smith, 2022) included 42 cases of follow-up and was discussed by the authors.'
            result=self.client.post(self.origin+'/api/rewrite',json={'text':text,'mode':'literature','level':'dba','length':'same'},headers=self.headers)
            self.assertEqual(result.status_code,200,result.text)
            self.assertIn('42 cases',result.json()['text'])
            kwargs=ctor.return_value.responses.create.call_args.kwargs
            self.assertFalse(kwargs['store'])
            self.assertEqual(ctor.call_args.kwargs['api_key'], os.environ['OPENAI_API_KEY'])


if __name__=='__main__':
    unittest.main()